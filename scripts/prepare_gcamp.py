"""Import published per-fly GCaMP6s traces without modifying the source workbook.

Requires openpyxl for this import step only. Simulation/fitting consume JSON.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
URL = 'https://cdn.elifesciences.org/articles/79887/elife-79887-fig3-data1-v3.xlsx'
SHA256 = '8258ac1e7bd49ff9dae01dca796d85dadbb595940a9a030abf20fe7d11ccc987'
# Left/right IDs independently cross-checked with Shiu 2024 Supplementary Table 1A
# and the pinned v783 annotations. These are cell-type matches across animals.
CELLS = {
    'Clavicle': ('20201014_SS48947_Clavicle', 'AN_GNG_30',
                 ['720575940655014049', '720575940632648868']),
    'G2N-1': ('20200804_SS47082_G2N-1', 'CB0616',
              ['720575940620874757', '720575940623718380']),
    'Zorro': ('20201019_SS67405_Zorro', 'CB0192',
              ['720575940629888530', '720575940611015122']),
}


def main():
    import openpyxl
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source', type=Path, default=ROOT/'outputs/gcamp-source/elife-79887-fig3-data1-v3.xlsx')
    p.add_argument('--output', type=Path, default=ROOT/'brain/calibration_data/shiu2022_taste.json')
    args = p.parse_args()
    if not args.source.exists():
        args.source.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(URL, timeout=60) as response:
            args.source.write_bytes(response.read())
    if hashlib.sha256(args.source.read_bytes()).hexdigest() != SHA256:
        raise ValueError('Source workbook checksum differs from the inspected publication version')
    with (ROOT/'brain/data/neuron_annotations.tsv').open() as f:
        annotations = {r['root_id']: r for r in csv.DictReader(f, delimiter='\t')}
    workbook = openpyxl.load_workbook(args.source, read_only=True, data_only=True)
    cells = []
    for name, (sheet, cell_type, ids) in CELLS.items():
        assert [annotations[i]['side'] for i in ids] == ['left', 'right']
        assert all(annotations[i]['cell_type'] == cell_type for i in ids)
        rows = list(workbook[sheet].iter_rows(values_only=True))
        assert rows[1][:3] == ('Frame', 'Time (seconds)', 'Taste stimulus')
        assert [r[0] for r in rows[2:]] == list(range(55))
        times = [r[1] for r in rows[2:]]
        stimulus = [r[2] for r in rows[2:]]
        assert all(abs(times[i]-1.2*i) < 1e-9 for i in range(55))
        assert [i for i, v in enumerate(stimulus) if v] == list(range(20, 26))
        traces = []
        for col, header in enumerate(rows[1][3:], start=3):
            if not header:
                continue
            fly, tastant = header.rsplit('_', 1)
            assert tastant in ('sugar', 'water', 'bitter')
            values = [r[col] for r in rows[2:]]
            # Two published G2N-1 sugar traces end early. Keep missing tail
            # samples null; never convert an unrecorded observation to zero.
            assert all(v is None or isinstance(v, (float, int)) for v in values)
            traces.append({'fly_id': fly, 'tastant': tastant,
                           'source_column': openpyxl.utils.get_column_letter(col+1),
                           'dff': values})
        flies = sorted({t['fly_id'] for t in traces})
        assert len(traces) == 3*len(flies)
        # A fixed chronological split, by whole fly, never by time point.
        held_out = flies[-max(2, len(flies)//3):]
        for trace in traces:
            trace['split'] = 'test' if trace['fly_id'] in held_out else 'train'
        cells.append({'name': name, 'cell_type': cell_type, 'flywire_ids': ids,
                      'sheet': sheet, 'description': rows[0][0],
                      'time_s': times, 'stimulus': stimulus, 'traces': traces})
    result = {
        'schema': 'gcamp-taste-pilot-v1',
        'source': {'paper': 'Shiu, Sterne et al., eLife 2022, Figure 3 source data 1',
                   'doi': '10.7554/eLife.79887', 'url': URL, 'sha256': SHA256,
                   'license': 'CC BY 4.0; attribution to the original authors',
                   'mapping_reference': 'Shiu et al. Nature 2024 Supplementary Table 1A; pinned v783 annotations'},
        'indicator': 'GCaMP6s', 'units': 'delta_F_over_F',
        'condition': 'food-deprived mated female flies; 1 M sucrose on the proboscis',
        'protocol': {'stimulus_on_s': 24.0, 'stimulus_off_s': 31.2,
                     'baseline_frames': [9, 18], 'evaluation_window_s': [18.0, 43.2],
                     'timing_basis': 'Workbook Time (seconds) and stimulus column; flags held until the next frame',
                     'timing_caveat': 'The workbook uses 1.2 s/frame; Methods describe nominal 0.8 Hz. Two-photon sheets have a larger timing discrepancy and are excluded.'},
        'matching': 'Cell-type correspondence across flies. Recording side is unspecified; model readout averages the listed bilateral homologues.',
        'missing_data': 'Original blank samples remain null; fitting masks them explicitly.',
        'split': 'Last third of chronological fly identifiers per cell type held out, minimum two flies. All tastants from a fly stay together.',
        'cells': cells,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, allow_nan=False, separators=(',', ':'))+'\n')
    print(json.dumps({'output': str(args.output), 'cells': len(cells),
                      'traces': sum(len(c['traces']) for c in cells),
                      'sugar_train': sum(t['split']=='train' for c in cells for t in c['traces'] if t['tastant']=='sugar'),
                      'sugar_test': sum(t['split']=='test' for c in cells for t in c['traces'] if t['tastant']=='sugar')}))


if __name__ == '__main__':
    main()
